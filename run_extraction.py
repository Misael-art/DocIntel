r"""
DocIntel - Stage 3: Production Extraction Runner
Self-contained single-file script. No relative imports needed.

Strategy:
  1. FASE_1 first (26k files - critical docs): hash + text extraction
  2. FASE_2 second (1.7M files): hash only for non-excluded paths
  3. Skip node_modules, .git, .venv, __pycache__ from deep analysis
  4. Batch commits every 5000 files to avoid locks
  5. Resume-safe: skips files already in 'HASH_CALCULADO' or 'EXTRAIDO' status

Run: python F:\DocIntel\run_extraction.py
"""
import os
import sys
import time
import sqlite3
import hashlib
import traceback
from datetime import datetime

# ─── CONFIG ────────────────────────────────────────────────────────────────────
DB_PATH  = r"F:\DocIntel\output\inventario_global.db"
LOG_PATH = r"F:\DocIntel\output\reports\extracao_log.md"
STATUS_PATH = r"F:\DocIntel\output\reports\status_execucao.md"

# Directories to skip for DEEP extraction (hashing + text)
SKIP_DIRS = {
    'node_modules', '.git', '.venv', '__pycache__',
    '.next', '.nuxt', 'dist', 'build', 'target',
    '.gradle', '.mvn', 'venv', 'env', '.tox',
}

# Extensions eligible for text extraction
TEXT_EXTS = {'.pdf', '.docx', '.doc', '.xlsx', '.xls', '.txt', '.md',
             '.rtf', '.odt', '.csv', '.json', '.xml', '.html', '.htm'}

BATCH_SIZE = 5000
# ─── END CONFIG ───────────────────────────────────────────────────────────────


def get_conn():
    conn = sqlite3.connect(DB_PATH, timeout=60)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.execute("PRAGMA cache_size=-128000")
    conn.row_factory = sqlite3.Row
    return conn


def should_skip(path: str) -> bool:
    """Returns True if path contains any excluded directory."""
    parts = path.replace('/', '\\').split('\\')
    return any(p.lower() in SKIP_DIRS for p in parts)


def sha256(path: str, block=65536) -> str | None:
    h = hashlib.sha256()
    try:
        with open(path, 'rb') as f:
            for chunk in iter(lambda: f.read(block), b''):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return None


def extract_text(path: str, ext: str) -> str:
    """Extracts text preview from document files."""
    try:
        if ext == '.pdf':
            import pypdf
            reader = pypdf.PdfReader(path, strict=False)
            text = '\n'.join(p.extract_text() or '' for p in reader.pages[:10])
            return text[:50000]
        elif ext in ('.docx',):
            import docx
            doc = docx.Document(path)
            return '\n'.join(p.text for p in doc.paragraphs)[:50000]
        elif ext in ('.xlsx', '.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = wb.active
            rows = []
            for r in sheet.iter_rows(max_row=50, values_only=True):
                rows.append(' | '.join(str(c) for c in r if c is not None))
            return '\n'.join(rows)[:20000]
        elif ext in ('.txt', '.md', '.csv', '.html', '.htm', '.xml', '.json'):
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read(20000)
    except Exception:
        pass
    return ''


def write_status(processed, total, skipped, errors, start_time, current_fase, current_file):
    elapsed = time.time() - start_time
    rate = processed / elapsed if elapsed > 0 else 0
    eta = (total - processed) / rate if rate > 0 else 0
    ts = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    content = f"""# Status de Execucao — DocIntel Pipeline

> Ultima atualizacao: {ts}

| Campo | Valor |
|-------|-------|
| **Fase Atual** | Etapa 3 — Extração ({current_fase}) |
| **Arquivos Processados** | {processed:,} / {total:,} |
| **Pulados (excluídos)** | {skipped:,} |
| **Erros** | {errors:,} |
| **Taxa de Processamento** | {rate:,.0f} arquivos/s |
| **Tempo Decorrido** | {elapsed:.0f}s ({elapsed/60:.1f} min) |
| **ETA** | {eta:.0f}s ({eta/60:.1f} min) |
| **Ultimo Arquivo** | `{current_file[-80:] if current_file else '—'}` |
| **Proxima Acao** | Continuar extração |
"""
    with open(STATUS_PATH, 'w', encoding='utf-8') as f:
        f.write(content)


def run():
    conn = get_conn()
    cursor = conn.cursor()

    # Count targets for FASE_1 and FASE_2 not yet processed
    cursor.execute("""
        SELECT COUNT(*) FROM files
        WHERE fase_correspondente IN ('FASE_1', 'FASE_2')
          AND (status_indexacao = 'PENDENTE' OR status_indexacao IS NULL)
    """)
    total = cursor.fetchone()[0]
    print(f"[EXTRACTION] Alvo: {total:,} arquivos críticos (FASE_1 + FASE_2) não processados.")

    if total == 0:
        print("[EXTRACTION] Nada a processar. Banco já atualizado.")
        conn.close()
        return

    # Fetch IDs + paths ordered by priority (FASE_1 first, then smallest files)
    cursor.execute("""
        SELECT id, caminho_completo, extensao, tamanho_bytes, fase_correspondente
        FROM files
        WHERE fase_correspondente IN ('FASE_1', 'FASE_2')
          AND (status_indexacao = 'PENDENTE' OR status_indexacao IS NULL)
        ORDER BY
            CASE fase_correspondente WHEN 'FASE_1' THEN 1 ELSE 2 END,
            tamanho_bytes ASC
    """)
    rows = cursor.fetchall()

    start_time = time.time()
    processed = 0
    skipped   = 0
    errors    = 0
    batch_updates = []
    last_path = ''

    print(f"[EXTRACTION] Iniciando às {datetime.now().strftime('%H:%M:%S')}...")

    for row in rows:
        fid     = row['id']
        path    = row['caminho_completo']
        ext     = (row['extensao'] or '').lower()
        fase    = row['fase_correspondente']
        last_path = path

        # --- Skip logic ---
        skip_reason = None
        if should_skip(path):
            # Identify which excluded dir was matched
            parts = path.replace('/', '\\').split('\\')
            matched = next((p for p in parts if p.lower() in SKIP_DIRS), 'excluded_dir')
            skip_reason = f'SKIP_RULE:{matched}'
            batch_updates.append(('EXCLUIDO', None, None, skip_reason, fid))
            skipped += 1
        elif not os.path.exists(path):
            skip_reason = 'AUSENTE:arquivo_nao_encontrado'
            batch_updates.append(('AUSENTE', None, None, skip_reason, fid))
            errors += 1
        else:
            # --- Hash ---
            doc_hash = sha256(path)

            # --- Text (only for relevant extensions and FASE_1) ---
            text = ''
            if ext in TEXT_EXTS and fase == 'FASE_1':
                text = extract_text(path, ext)

            new_status = 'EXTRAIDO' if text else 'HASH_CALCULADO'
            batch_updates.append((new_status, doc_hash, text or None, None, fid))
            processed += 1

        # --- Batch commit ---
        if len(batch_updates) >= BATCH_SIZE:
            cursor.executemany("""
                UPDATE files SET
                    status_indexacao = ?,
                    hash_sha256 = ?,
                    texto_extraido = ?,
                    observacoes = COALESCE(?, observacoes),
                    updated_at = datetime('now')
                WHERE id = ?
            """, batch_updates)
            conn.commit()
            batch_updates.clear()

            # Progress report
            total_done = processed + skipped + errors
            write_status(total_done, total, skipped, errors, start_time, fase, last_path)
            rate = total_done / (time.time() - start_time)
            eta  = (total - total_done) / rate if rate > 0 else 0
            print(f"  [{total_done:,}/{total:,}] {rate:,.0f} arq/s | ETA: {eta/60:.1f} min | Erros: {errors}")

    # Final flush
    if batch_updates:
        cursor.executemany("""
            UPDATE files SET
                status_indexacao = ?,
                hash_sha256 = ?,
                texto_extraido = ?,
                observacoes = COALESCE(?, observacoes),
                updated_at = datetime('now')
            WHERE id = ?
        """, batch_updates)
        conn.commit()

    total_done = processed + skipped + errors
    final_fase = 'CONCLUIDO' if total_done >= total else (fase if rows else 'SEM_DADOS')
    write_status(total_done, total, skipped, errors, start_time, final_fase, last_path)

    conn.close()

    elapsed = time.time() - start_time
    summary = (f"Processados: {processed:,} | Pulados: {skipped:,} | "
               f"Erros: {errors:,} | Tempo: {elapsed:.0f}s ({elapsed/60:.1f} min)")
    print(f"\n[EXTRACTION] CONCLUÍDO. {summary}")

    # Write final log
    with open(LOG_PATH, 'a', encoding='utf-8') as f:
        f.write(f"\n## Execução {datetime.now().isoformat()}\n{summary}\n")


if __name__ == "__main__":
    run()
